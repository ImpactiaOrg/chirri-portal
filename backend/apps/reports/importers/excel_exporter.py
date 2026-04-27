"""Exporter xlsx de un Report existente (DEV-83 · Etapa 1).

Genera un xlsx con la misma shape que el template vacío de `excel_writer`
pero poblado con la data de un Report real. Uso principal: producir el
`reporte-abril-ejemplo.xlsx` que sirve como few-shot canónico (humano y
LLM) y referencia ante dudas de formato.
"""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from pathlib import PurePosixPath

from openpyxl.worksheet.worksheet import Worksheet

from apps.reports.models import (
    ChartBlock,
    ImageBlock,
    KpiGridBlock,
    Report,
    TableBlock,
    TextImageBlock,
    TopContentsBlock,
    TopCreatorsBlock,
)

from . import schema as s
from .excel_writer import build_skeleton, to_bytes


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------
def export(report: Report) -> BytesIO:
    """Serializa `report` al mismo xlsx que produciría el importer."""
    wb = build_skeleton()
    names = _assign_names(report)

    _populate_reporte(wb[s.SHEET_REPORTE], report, names)
    _populate_textimage(wb[s.SHEET_TEXTIMAGE], report, names)
    _populate_imagenes(wb[s.SHEET_IMAGENES], report, names)
    _populate_kpis(wb[s.SHEET_KPIS], report, names)
    _populate_tables(wb[s.SHEET_TABLES], report, names)
    _populate_topcontents(wb[s.SHEET_TOPCONTENTS], report, names)
    _populate_topcreators(wb[s.SHEET_TOPCREATORS], report, names)
    _populate_charts(wb[s.SHEET_CHARTS], report, names)

    return to_bytes(wb)


# ---------------------------------------------------------------------------
# Name assignment: block.pk → `{prefix}_{N}` (ej. textimage_1, kpi_1, imagen_2)
# ---------------------------------------------------------------------------
def _assign_names(report: Report) -> dict[int, str]:
    counters: dict[str, int] = {}
    names: dict[int, str] = {}
    for block in report.blocks.all().order_by("order"):
        type_name = type(block).__name__
        prefix = s.TYPE_PREFIX[type_name]
        counters[prefix] = counters.get(prefix, 0) + 1
        names[block.pk] = f"{prefix}_{counters[prefix]}"
    return names


# ---------------------------------------------------------------------------
# Per-sheet populators
# ---------------------------------------------------------------------------
def _populate_reporte(
    ws: Worksheet, report: Report, names: dict[int, str]
) -> None:
    kv_values = {
        "tipo": s.KIND_LABELS.get(report.kind, report.kind),
        "fecha_inicio": _fmt_date(report.period_start),
        "fecha_fin": _fmt_date(report.period_end),
        "titulo": report.title,
        "intro": report.intro_text,
        "conclusiones": report.conclusions_text,
    }
    for idx, (key, _type, _required, _example) in enumerate(s.REPORTE_KV_ROWS, start=2):
        ws.cell(row=idx, column=2, value=kv_values.get(key, ""))

    # Layout rows — encontrar fila del header "orden" y empezar debajo.
    layout_row = _find_row_starting_with(ws, "orden")
    if layout_row is None:
        return
    for i, block in enumerate(report.blocks.all().order_by("order"), start=1):
        row = layout_row + i
        ws.cell(row=row, column=1, value=block.order)
        ws.cell(row=row, column=2, value=names[block.pk])


def _populate_textimage(
    ws: Worksheet, report: Report, names: dict[int, str]
) -> None:
    for i, block in enumerate(
        TextImageBlock.objects.filter(report=report).order_by("order"), start=2
    ):
        _write_row(ws, i, s.TEXTIMAGE_HEADERS, {
            "nombre": names[block.pk],
            "title": block.title,
            "body": block.body,
            "imagen": _filename(block.image),
            "image_alt": block.image_alt,
            "image_position": block.image_position,
            "columns": block.columns,
        })


def _populate_imagenes(
    ws: Worksheet, report: Report, names: dict[int, str]
) -> None:
    for i, block in enumerate(
        ImageBlock.objects.filter(report=report).order_by("order"), start=2
    ):
        _write_row(ws, i, s.IMAGENES_HEADERS, {
            "nombre": names[block.pk],
            "title": block.title,
            "caption": block.caption,
            "imagen": _filename(block.image),
            "image_alt": block.image_alt,
        })


def _populate_kpis(
    ws: Worksheet, report: Report, names: dict[int, str]
) -> None:
    row = 2
    for block in KpiGridBlock.objects.filter(report=report).order_by("order"):
        for tile in block.tiles.all().order_by("order"):
            _write_row(ws, row, s.KPIS_HEADERS, {
                "nombre": names[block.pk],
                "block_title": block.title,
                "item_orden": tile.order,
                "label": tile.label,
                "value": _num(tile.value),
                "period_comparison": _num(tile.period_comparison),
            })
            row += 1


def _populate_tables(
    ws: Worksheet, report: Report, names: dict[int, str]
) -> None:
    row = 2
    for block in TableBlock.objects.filter(report=report).order_by("order"):
        for r in block.rows.all().order_by("order"):
            cells_padded = list(r.cells) + [""] * (8 - len(r.cells))
            values = {
                "nombre": names[block.pk],
                "block_title": block.title,
                "block_show_total": "TRUE" if block.show_total else "FALSE",
                "row_orden": r.order,
                "is_header": "TRUE" if r.is_header else "FALSE",
            }
            for i, cell in enumerate(cells_padded[:8], start=1):
                values[f"cell_{i}"] = cell
            _write_row(ws, row, s.TABLES_HEADERS, values)
            row += 1


def _populate_topcontents(
    ws: Worksheet, report: Report, names: dict[int, str]
) -> None:
    row = 2
    for block in TopContentsBlock.objects.filter(report=report).order_by("order"):
        for item in block.items.all().order_by("order"):
            _write_row(ws, row, s.TOPCONTENTS_HEADERS, {
                "nombre": names[block.pk],
                "block_title": block.title,
                "block_network": s.NETWORK_LABELS.get(block.network, ""),
                "block_period_label": block.period_label,
                "block_limit": block.limit,
                "item_orden": item.order,
                "imagen": _filename(item.thumbnail),
                "caption": item.caption,
                "post_url": item.post_url,
                "source_type": s.SOURCE_TYPE_LABELS.get(item.source_type, ""),
                "views": item.views,
                "likes": item.likes,
                "comments": item.comments,
                "shares": item.shares,
                "saves": item.saves,
            })
            row += 1


def _populate_topcreators(
    ws: Worksheet, report: Report, names: dict[int, str]
) -> None:
    row = 2
    for block in TopCreatorsBlock.objects.filter(report=report).order_by("order"):
        for item in block.items.all().order_by("order"):
            _write_row(ws, row, s.TOPCREATORS_HEADERS, {
                "nombre": names[block.pk],
                "block_title": block.title,
                "block_network": s.NETWORK_LABELS.get(block.network, ""),
                "block_period_label": block.period_label,
                "block_limit": block.limit,
                "item_orden": item.order,
                "imagen": _filename(item.thumbnail),
                "handle": item.handle,
                "post_url": item.post_url,
                "views": item.views,
                "likes": item.likes,
                "comments": item.comments,
                "shares": item.shares,
            })
            row += 1


def _populate_charts(
    ws: Worksheet, report: Report, names: dict[int, str]
) -> None:
    row = 2
    for block in ChartBlock.objects.filter(report=report).order_by("order"):
        for p in block.data_points.all().order_by("order"):
            _write_row(ws, row, s.CHARTS_HEADERS, {
                "nombre": names[block.pk],
                "block_title": block.title,
                "block_network": s.NETWORK_LABELS.get(block.network, ""),
                "chart_type": block.chart_type,
                "point_orden": p.order,
                "point_label": p.label,
                "point_value": _num(p.value),
            })
            row += 1


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _write_row(
    ws: Worksheet, row: int, headers: list[str], values: dict,
) -> None:
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=row, column=col_idx, value=values.get(header, ""))


def _filename(imagefield) -> str:
    """Extrae solo el nombre del archivo de un ImageField (sin path)."""
    if not imagefield or not imagefield.name:
        return ""
    return PurePosixPath(imagefield.name).name


def _num(value):
    """Normaliza Decimal/None/0 para Excel: None → '', Decimal → float."""
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    return value


def _fmt_date(d) -> str:
    if d is None:
        return ""
    return d.strftime("%d/%m/%Y")


def _find_row_starting_with(ws: Worksheet, header: str) -> int | None:
    for row in range(1, ws.max_row + 1):
        cell = ws.cell(row=row, column=1).value
        if cell == header:
            return row
    return None
